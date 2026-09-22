#!/usr/bin/env python3
"""Write a week of jobs into the RS MS 2026 master schedule.

Each job block goes in its lead crew member's column, starting on the row whose
time in column A matches the job's start (anything before 10am sits on the first
row, as in the hand-filled weeks).
"""
import argparse
import importlib.util
from copy import copy

import openpyxl


def load(path):
    spec = importlib.util.spec_from_file_location("week", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def time_rows(ws):
    rows = {}
    for r in range(9, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v:
            rows[int(v)] = r
    return rows


def row_for(rows, hhmm):
    """First grid row at or after hhmm; earlier starts sit on the first row."""
    later = [t for t in sorted(rows) if t >= hhmm]
    return rows[later[0]] if later else rows[max(rows)]


def day_columns(ws, date_label):
    for c in range(3, ws.max_column + 1):
        if str(ws.cell(row=6, column=c).value).strip() == date_label:
            return {str(ws.cell(row=8, column=cc).value).strip(): cc
                    for cc in range(c, c + 8)
                    if ws.cell(row=8, column=cc).value}
    raise SystemExit(f"date column not found: {date_label}")


def fill(week, path, out):
    wb = openpyxl.load_workbook(path)
    ws = wb[week.SHEET]
    rows = time_rows(ws)
    style_src = None
    written = []

    for date_label, crew, start, lines in week.PLAN:
        cols = day_columns(ws, date_label)
        if crew not in cols:
            raise SystemExit(f"{crew} has no column on {date_label}")
        col = cols[crew]
        r0 = row_for(rows, start)
        if style_src is None:
            style_src = ws.cell(row=9, column=5)._style
        for i, text in enumerate(lines):
            cell = ws.cell(row=r0 + i, column=col)
            if cell.value not in (None, ""):
                raise SystemExit(f"{date_label} {crew} row {r0 + i} already holds {cell.value!r}")
            cell.value = text
            cell._style = copy(style_src)
        written.append((date_label, crew, r0, lines[0]))

    wb.save(out)
    return written


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("week")
    ap.add_argument("--master", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    for d, crew, r, head in fill(load(args.week), args.master, args.out):
        print(f"{d:22} {crew:8} row {r:<3} {head}")


if __name__ == "__main__":
    main()
