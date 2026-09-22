#!/usr/bin/env python3
"""Build the weekly crew schedule .docx from a jobs file, reusing an existing
schedule as the template so fonts, margins and the layout stay identical.

    python3 schedule/make_schedule.py jobs_2026-09-28.py --template last_week.docx --out out.docx
"""
import argparse
import importlib.util
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

FONT, SIZE = "Arial", Pt(10)


def load(path):
    spec = importlib.util.spec_from_file_location("jobs", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def line(doc, text="", bold=False, underline=False, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = SIZE
    run.bold = bold or None
    run.underline = underline or None
    return p


def block(doc, heading, lines):
    line(doc, heading)
    for t in lines:
        line(doc, t)


def build(jobs, template, out):
    doc = Document(template)
    clear_body(doc)

    by_person = {}
    for day, key, heading, lines, crew in jobs.JOBS:
        if crew:
            by_person.setdefault(crew, {}).setdefault(day, []).append((key, heading, lines))
    unassigned = {}
    for day, key, heading, lines, crew in jobs.JOBS:
        if not crew:
            unassigned.setdefault(day, []).append((key, heading, lines))

    for person in jobs.ROSTER:
        line(doc, f"{person}’s Schedule", bold=True, underline=True, center=True)
        line(doc)
        line(doc)
        for day, label in jobs.DAYS:
            line(doc, label)
            entries = sorted(by_person.get(person, {}).get(day, []))
            if person == "Evelyn" and not entries:
                line(doc, "Clear Video")
            elif not entries:
                line(doc, "To Be Assigned")
            else:
                for i, (_, heading, lines) in enumerate(entries):
                    if i:
                        line(doc)
                    block(doc, heading, lines)
            line(doc)
        line(doc)
        line(doc)

    line(doc, "Jobs To Be Assigned", bold=True, underline=True, center=True)
    line(doc)
    line(doc)
    for day, label in jobs.DAYS:
        line(doc, label)
        entries = sorted(unassigned.get(day, []))
        if not entries:
            line(doc, "Nil")
        for i, (_, heading, lines) in enumerate(entries):
            if i:
                line(doc)
            block(doc, heading, lines)
        line(doc)
    line(doc)
    line(doc)

    line(doc, "Reminders", bold=True, underline=True, center=True)
    line(doc)
    line(doc)
    for heading, lines in jobs.NOTES:
        block(doc, heading, lines)
        line(doc)

    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("jobs")
    ap.add_argument("--template", required=True, help="last week's schedule .docx")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    path = build(load(args.jobs), args.template, args.out)
    print(path, os.path.getsize(path), "bytes")


if __name__ == "__main__":
    main()
